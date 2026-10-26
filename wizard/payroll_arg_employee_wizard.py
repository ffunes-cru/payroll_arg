"""@file
Alta y actualización masiva de empleados desde Excel / CSV, y exportación de la plantilla o de los
empleados actuales con el mismo formato.
"""
import base64
import csv
import datetime
import io
import openpyxl
import xlsxwriter
from odoo import api, fields, models
from odoo.exceptions import UserError



class PayrollArgEmployeeWizard(models.TransientModel):
    """Asistente de importación / exportación de empleados.

    Columnas: legajo, nombre, CUIL, tipo de salario, código CCT, código de categoría, sueldo o valor
    hora, fecha de ingreso, código de obra social, forma de pago, CBU y lugar de trabajo.
    """

    _name = "payroll.arg.employee.wizard"
    _description = "Importador y Exportador Masivo de Empleados"

    operation_mode = fields.Selection(
        [
            ("export", "Exportar Plantilla / Empleados"),
            ("import", "Importar Empleados desde Archivo"),
        ],
        string="Operación a Realizar",
        default="export",
        required=True,
    )
    export_type = fields.Selection(
        [
            ("template", "Plantilla Base Vacía (con filas de ejemplo)"),
            ("all_employees", "Exportar Todos los Empleados Actuales"),
        ],
        string="Contenido de Exportación",
        default="all_employees",
    )
    data_file = fields.Binary(string="Archivo a Importar (Excel .xlsx o CSV)")
    filename = fields.Char(string="Nombre del Archivo")
    export_file = fields.Binary(string="Archivo Descargable", readonly=True)
    export_filename = fields.Char(string="Nombre Archivo Exportado", readonly=True)
    state = fields.Selection(
        [
            ("choose", "Configuración"),
            ("done", "Completado"),
        ],
        default="choose",
    )
    result_summary = fields.Text(string="Resultado de la Operación", readonly=True)

    def action_export_excel(self):
        """Genera el Excel: plantilla con dos filas de ejemplo o todos los empleados actuales.

        @return Acción que reabre el asistente con el archivo para descargar.
        """
        self.ensure_one()
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        worksheet = workbook.add_worksheet("Empleados")

        # Estilos
        header_format = workbook.add_format({
            "bold": True,
            "font_color": "white",
            "bg_color": "#1B365D",
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "text_wrap": True,
        })
        cell_format = workbook.add_format({"border": 1, "valign": "vcenter"})
        number_format = workbook.add_format({"border": 1, "num_format": "#,##0.00", "align": "right"})
        date_format = workbook.add_format({"border": 1, "num_format": "yyyy-mm-dd", "align": "center"})
        code_format = workbook.add_format({"border": 1, "align": "center"})

        headers = [
            ("Legajo", 12),
            ("Nombre y Apellido", 30),
            ("CUIL", 16),
            ("Tipo Salario", 16),
            ("Código Convenio CCT", 20),
            ("Código Categoría", 18),
            ("Sueldo / Valor Hora", 20),
            ("Fecha Ingreso", 15),
            ("Código Obra Social", 18),
            ("Forma Pago", 14),
            ("CBU", 26),
            ("Lugar de Trabajo", 20),
        ]

        for col_idx, (header_name, col_width) in enumerate(headers):
            worksheet.write(0, col_idx, header_name, header_format)
            worksheet.set_column(col_idx, col_idx, col_width)

        row_idx = 1
        if self.export_type == "template":
            # Filas de ejemplo explicativas
            samples = [
                (
                    "1001",
                    "Pérez Juan Carlos",
                    "20-33445566-9",
                    "Jornalizado",
                    "260/75",
                    "OF",
                    5400.0,
                    "2021-03-15",
                    "126205",
                    "3",
                    "0720262188000036085521",
                    "Córdoba",
                ),
                (
                    "2001",
                    "Gómez Marcos Antonio",
                    "20-28112233-4",
                    "Mensualizado",
                    "260/75",
                    "ADM1",
                    850000.0,
                    "2019-06-01",
                    "126205",
                    "3",
                    "0170099220000067891234",
                    "Córdoba",
                ),
            ]
            for sample in samples:
                worksheet.write(row_idx, 0, sample[0], code_format)
                worksheet.write(row_idx, 1, sample[1], cell_format)
                worksheet.write(row_idx, 2, sample[2], code_format)
                worksheet.write(row_idx, 3, sample[3], code_format)
                worksheet.write(row_idx, 4, sample[4], code_format)
                worksheet.write(row_idx, 5, sample[5], code_format)
                worksheet.write(row_idx, 6, sample[6], number_format)
                worksheet.write(row_idx, 7, sample[7], date_format)
                worksheet.write(row_idx, 8, sample[8], code_format)
                worksheet.write(row_idx, 9, sample[9], code_format)
                worksheet.write(row_idx, 10, sample[10], code_format)
                worksheet.write(row_idx, 11, sample[11], cell_format)
                row_idx += 1
        else:
            # Exportar todos los empleados actuales
            employees = self.env["hr.employee"].search([], order="name")
            for emp in employees:
                contract = emp.current_version_id or emp.version_id
                wage_type_str = "Jornalizado" if (contract and contract.wage_type == "hourly") else "Mensualizado"
                cct_code = (contract and contract.cct_id and contract.cct_id.code) or (emp.cct_id and emp.cct_id.code) or ""
                cat_code = (contract and contract.labor_category_id and contract.labor_category_id.code) or (emp.labor_category_id and emp.labor_category_id.code) or ""
                wage_val = contract.wage if contract else 0.0
                date_start_str = str(contract.date_start or "") if contract and contract.date_start else ""
                os_code = emp.health_insurance_id.code if emp.health_insurance_id else ""

                worksheet.write(row_idx, 0, str(emp.legajo or emp.id), code_format)
                worksheet.write(row_idx, 1, emp.name or "", cell_format)
                worksheet.write(row_idx, 2, emp.cuil or "", code_format)
                worksheet.write(row_idx, 3, wage_type_str, code_format)
                worksheet.write(row_idx, 4, cct_code, code_format)
                worksheet.write(row_idx, 5, cat_code, code_format)
                worksheet.write(row_idx, 6, wage_val, number_format)
                worksheet.write(row_idx, 7, date_start_str, date_format)
                worksheet.write(row_idx, 8, os_code, code_format)
                worksheet.write(row_idx, 9, emp.payment_type or "3", code_format)
                worksheet.write(row_idx, 10, emp.cbu or "", code_format)
                worksheet.write(row_idx, 11, emp.workplace_location or "Córdoba", cell_format)
                row_idx += 1

        workbook.close()
        output.seek(0)
        file_data = base64.b64encode(output.read())
        filename = "plantilla_empleados_nomina_arg.xlsx" if self.export_type == "template" else "empleados_nomina_arg.xlsx"

        self.write({
            "export_file": file_data,
            "export_filename": filename,
            "state": "done",
            "result_summary": f"Archivo generado exitosamente ({row_idx - 1} registros). Puede descargarlo a continuación.",
        })
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "views": [[False, "form"]],
            "target": "new",
        }

    @api.model
    def _parse_date_value(self, date_val):
        """Fecha de una celda: acepta fechas de Excel y texto AAAA-MM-DD, DD/MM/AAAA, DD-MM-AAAA y
        AAAA/MM/DD.

        @param date_val Valor de la celda.
        @return date, o False si no se puede leer.
        """
        if not date_val:
            return False
        if isinstance(date_val, datetime.date):
            return date_val
        if isinstance(date_val, datetime.datetime):
            return date_val.date()
        val_str = str(date_val).strip()
        if not val_str:
            return False
        if " " in val_str:
            val_str = val_str.split(" ")[0]
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.datetime.strptime(val_str, fmt).date()
            except ValueError:
                continue
        try:
            return fields.Date.to_date(val_str)
        except Exception:
            return False

    def action_import_file(self):
        """Crea o actualiza empleados desde el archivo.

        Busca al empleado por CUIL y luego por legajo. Si la categoría tiene tipo, manda sobre la columna
        de tipo de salario; sin sueldo usa el básico de la categoría. Asigna calendario, estructura y
        periodicidad por defecto del tipo de salario. Los errores de una fila no frenan las demás.

        @return Acción que reabre el asistente con el resumen.
        @throws UserError Sin archivo, Excel ilegible o sin filas.
        """
        self.ensure_one()
        if not self.data_file:
            raise UserError("Por favor seleccione un archivo Excel (.xlsx) o CSV para importar.")

        file_bytes = base64.b64decode(self.data_file)
        rows_to_process = []

        is_excel = (self.filename and self.filename.lower().endswith(".xlsx")) or file_bytes.startswith(b"PK")
        if is_excel:
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                ws = wb.active
                for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if r_idx == 0:
                        continue  # Saltear cabecera
                    if not row or not any(row):
                        continue
                    rows_to_process.append([str(c or "").strip() for c in row])
            except Exception as e:
                raise UserError(f"Error al leer archivo Excel: {e}")
        else:
            try:
                content = file_bytes.decode("utf-8-sig")
            except Exception:
                content = file_bytes.decode("latin1")
            delimiter = ";" if ";" in content.splitlines()[0] else ","
            reader = csv.reader(io.StringIO(content), delimiter=delimiter)
            for r_idx, row in enumerate(reader):
                if r_idx == 0:
                    continue
                if not row or not any(row):
                    continue
                rows_to_process.append([str(c).strip() for c in row])

        if not rows_to_process:
            raise UserError("El archivo no contiene filas de datos para procesar.")

        created_count = 0
        updated_count = 0
        errors = []

        cct_model = self.env["afip.cct"]
        cat_model = self.env["afip.labor.category"]
        os_model = self.env["afip.health.insurance"]
        emp_model = self.env["hr.employee"]

        for idx, row in enumerate(rows_to_process, start=2):
            try:
                # Columnas mínimas esperadas: legajo, nombre, cuil
                legajo = row[0] if len(row) > 0 else ""
                name = row[1] if len(row) > 1 else ""
                cuil_raw = row[2] if len(row) > 2 else ""
                wage_type_raw = (row[3] or "").strip().lower() if len(row) > 3 else "mensualizado"
                cct_code = row[4] if len(row) > 4 else ""
                cat_code = row[5] if len(row) > 5 else ""
                wage_raw = row[6] if len(row) > 6 else ""
                date_start_raw = row[7] if len(row) > 7 else ""
                os_code = row[8] if len(row) > 8 else ""
                payment_type = row[9] if len(row) > 9 else "3"
                cbu = row[10] if len(row) > 10 else ""
                workplace = row[11] if len(row) > 11 else "Córdoba"

                if not name and not cuil_raw:
                    continue

                cuil_clean = "".join([c for c in cuil_raw if c.isdigit()])
                cbu_clean = "".join([c for c in cbu if c.isdigit()])

                # Buscar CCT
                cct = cct_model.search([("code", "=", cct_code)], limit=1) if cct_code else False
                if not cct and cct_code:
                    cct = cct_model.search([("name", "ilike", cct_code)], limit=1)

                # Buscar Categoría
                cat = cat_model.search([("code", "=", cat_code)], limit=1) if cat_code else False
                if not cat and cat_code:
                    cat = cat_model.search([("name", "ilike", cat_code)], limit=1)

                # Determinar wage_type:
                # Si la categoría tiene definido category_type ('hourly' o 'monthly'), manda la categoría.
                # De lo contrario, se analiza lo escrito en la columna tipo de salario.
                if cat and cat.category_type:
                    wage_type = cat.category_type
                elif wage_type_raw and ("jornal" in wage_type_raw or "hora" in wage_type_raw):
                    wage_type = "hourly"
                else:
                    wage_type = "monthly"

                # Salario
                wage_val = 0.0
                if wage_raw:
                    try:
                        clean_num = str(wage_raw).replace("$", "").strip()
                        if "." in clean_num and "," in clean_num:
                            if clean_num.find(".") < clean_num.find(","):
                                clean_num = clean_num.replace(".", "").replace(",", ".")
                            else:
                                clean_num = clean_num.replace(",", "")
                        elif "," in clean_num:
                            clean_num = clean_num.replace(",", ".")
                        wage_val = float(clean_num)
                    except Exception:
                        wage_val = 0.0

                if wage_val == 0.0 and cat:
                    wage_val = cat.hourly_wage if wage_type == "hourly" else cat.monthly_wage

                # Obra Social
                os_rec = os_model.search([("code", "=", os_code)], limit=1) if os_code else False

                # Buscar si el empleado ya existe
                employee = False
                if cuil_clean:
                    employee = emp_model.search([("cuil", "like", cuil_clean)], limit=1)
                if not employee and legajo:
                    employee = emp_model.search([("legajo", "=", legajo)], limit=1)

                # Fecha ingreso / inicio contrato
                start_date = self._parse_date_value(date_start_raw)

                # Calendario, Estructura y Periodicidad
                cal = self.env["hr.version"]._get_default_calendar(wage_type)
                struct = self.env["hr.version"]._get_default_struct(wage_type)
                schedule_pay = "bi-weekly" if wage_type == "hourly" else "monthly"

                emp_vals = {
                    "name": name,
                    "cuil": cuil_clean or cuil_raw,
                    "legajo": legajo,
                    "payment_type": payment_type if payment_type in ["1", "2", "3"] else "3",
                    "cbu": cbu_clean or cbu,
                    "workplace_location": workplace or "Córdoba",
                    "wage_type": wage_type,
                    "schedule_pay": schedule_pay,
                    "wage": wage_val,
                }
                if struct:
                    emp_vals["struct_id"] = struct.id
                if cal:
                    emp_vals["resource_calendar_id"] = cal.id
                if cct:
                    emp_vals["cct_id"] = cct.id
                if cat:
                    emp_vals["labor_category_id"] = cat.id
                if os_rec:
                    emp_vals["health_insurance_id"] = os_rec.id
                if start_date:
                    emp_vals["contract_date_start"] = start_date
                    emp_vals["date_version"] = start_date

                if employee:
                    employee.write(emp_vals)
                    contract = employee.current_version_id or employee.version_id
                    if contract and (not contract.name or contract.name == "/"):
                        contract.name = f"Contrato - {employee.name}"
                    updated_count += 1
                else:
                    new_emp = emp_model.create(emp_vals)
                    contract = new_emp.current_version_id or new_emp.version_id
                    if contract and (not contract.name or contract.name == "/"):
                        contract.name = f"Contrato - {new_emp.name}"
                    created_count += 1

            except Exception as e:
                err_msg = str(e)
                if hasattr(e, "__context__") and e.__context__:
                    err_msg = f"{str(e.__context__)} (Detalle: {err_msg})"
                errors.append(f"Fila {idx} ({row[1] if len(row) > 1 else 'Sin nombre'}): {err_msg}")

        summary_lines = [
            f"Proceso de importación finalizado con éxito.",
            f"• Empleados Creados: {created_count}",
            f"• Empleados Actualizados: {updated_count}",
        ]
        if errors:
            summary_lines.append(f"• Errores ({len(errors)}):")
            summary_lines.extend([f"  - {err}" for err in errors[:50]])

        self.write({
            "state": "done",
            "result_summary": "\n".join(summary_lines),
        })
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "views": [[False, "form"]],
            "target": "new",
        }
